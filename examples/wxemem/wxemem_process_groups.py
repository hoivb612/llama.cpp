"""Generate group composition and change-waterfall slides using a supplied rules file."""

import argparse
from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
import sys
import textwrap

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from wxemem_compare import MIB, short_text
from wxemem_process import Row, aggregate, compare_snapshots, load_snapshot, mib

PALETTE = ("4263EB", "12A594", "EEA62B", "9362CB", "DC7188", "397C9C")
UNCLASSIFIED = "Unclassified"
UNKNOWN = "Unknown images"
AMBIGUOUS = "Ambiguous matches"
ROLLUP = "Other configured groups"
RESERVED = {n.lower() for n in (UNCLASSIFIED, UNKNOWN, AMBIGUOUS, ROLLUP)}
LOWER_ASCII = str.maketrans("ABCDEFGHIJKLMNOPQRSTUVWXYZ", "abcdefghijklmnopqrstuvwxyz")


@dataclass(frozen=True)
class Group:
    name: str
    patterns: tuple[str, ...]


@dataclass(frozen=True)
class Rules:
    path: Path
    digest: str
    text: str
    groups: tuple[Group, ...]


@dataclass(frozen=True)
class Bucket:
    name: str
    color: str
    rows: tuple[Row, ...]
    members: tuple[str, ...]
    configured: bool

    @property
    def a(self):
        return sum(r.a for r in self.rows)

    @property
    def b(self):
        return sum(r.b for r in self.rows)

    @property
    def delta(self):
        return self.b - self.a


def load_groups(path):
    path = Path(path).resolve(strict=True)
    raw = path.read_bytes()
    source = raw.decode("utf-8-sig")
    definitions = []
    names = set()
    for number, line in enumerate(source.splitlines(), 1):
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        prefix, colon, name = line.partition(":")
        if colon and prefix.strip().translate(LOWER_ASCII).startswith("group"):
            name = short_text(name.strip(), 90, f"Group name at line {number}")
            if name.lower() in names or name.lower() in RESERVED:
                raise ValueError(f"Duplicate or reserved group name at line {number}: {name}")
            names.add(name.lower())
            definitions.append((name, []))
        else:
            if not definitions:
                raise ValueError(f"Pattern before first group at {path}:{number}")
            if any(ord(c) < 32 for c in line):
                raise ValueError(f"Invalid control character in pattern at line {number}")
            pattern = line.translate(LOWER_ASCII)
            if pattern not in definitions[-1][1]:
                definitions[-1][1].append(pattern)
    if not definitions:
        raise ValueError("Grouping file contains no groups")
    for name, patterns in definitions:
        if not patterns:
            raise ValueError(f"Group has no patterns: {name}")
    return Rules(path, hashlib.sha256(raw).hexdigest(), source,
                 tuple(Group(n, tuple(p)) for n, p in definitions))


def classify(a, b, rules, metric="ws"):
    comparison = compare_snapshots(a, b, metric)
    aa, bb = aggregate(a, comparison.use_ws), aggregate(b, comparison.use_ws)
    classified = {g.name: [] for g in rules.groups}
    classified.update({UNCLASSIFIED: [], UNKNOWN: [], AMBIGUOUS: []})
    matches = {}
    for name in sorted(aa.keys() | bb.keys()):
        lower = name.translate(LOWER_ASCII)
        hits = tuple(g.name for g in rules.groups if any(p in lower for p in g.patterns))
        # Unknown image identities cannot be safely assigned by a substring rule.
        target = (UNKNOWN if lower.strip() == "<unknown>" else
                  hits[0] if len(hits) == 1 else AMBIGUOUS if hits else UNCLASSIFIED)
        amount_a, count_a = aa.get(name, (0, 0))
        amount_b, count_b = bb.get(name, (0, 0))
        classified[target].append(Row(name, amount_a, amount_b, count_a, count_b))
        matches[name] = hits
    buckets = [
        Bucket(g.name, PALETTE[i % len(PALETTE)], tuple(classified[g.name]), (g.name,), True)
        for i, g in enumerate(rules.groups)
    ]
    buckets.extend(Bucket(n, color, tuple(classified[n]), (), False)
                   for n, color in ((UNCLASSIFIED, "A4B6CE"), (UNKNOWN, "617389"), (AMBIGUOUS, "D9695C")))
    if (sum(v.a for v in buckets) != comparison.total_a
            or sum(v.b for v in buckets) != comparison.total_b
            or sum(len(v.rows) for v in buckets) != len(aa.keys() | bb.keys())):
        raise ValueError("Exclusive group assignment does not reconcile")
    return comparison, tuple(buckets), matches


def visible_buckets(buckets):
    configured = [v for v in buckets if v.configured]
    if len(configured) > 6:
        largest = sorted(configured, key=lambda v: (-max(v.a, v.b), v.name))[:5]
        names = {v.name for v in largest}
        rest = [v for v in configured if v.name not in names]
        configured = [v for v in configured if v.name in names]
        configured.append(Bucket(ROLLUP, "9362CB", tuple(r for v in rest for r in v.rows),
                                 tuple(v.name for v in rest), True))
        configured = [Bucket(v.name, PALETTE[i], v.rows, v.members, True)
                      for i, v in enumerate(configured)]
    return tuple(configured + [v for v in buckets if not v.configured and (v.rows or v.name == UNCLASSIFIED)])


def waterfall_steps(total_a, buckets, total_b):
    steps = [("A total", 0, total_a, total_a, "182A43")]
    current = total_a
    for v in buckets:
        following = current + v.delta
        if following < 0:
            raise ValueError("Waterfall intermediate total is negative")
        steps.append((v.name, min(current, following), max(current, following), v.delta, v.color))
        current = following
    if current != total_b:
        raise ValueError("Waterfall does not end at B")
    steps.append(("B total", 0, total_b, total_b, "182A43"))
    return tuple(steps)


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("F7F9FC")
    return slide


def rect(slide, x, y, w, h, color, name=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    return shape


def text(slide, x, y, w, h, value, size=12, color="182A43", bold=False, align=PP_ALIGN.LEFT, name=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, row in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text, p.alignment = row, align
        p.font.name, p.font.size, p.font.bold = "Aptos", Pt(size), bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(0)
    return shape


def line(slide, x1, y1, x2, y2, color="DCE3ED"):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = RGBColor.from_string(color)
    shape.line.width = Pt(1)


def header(slide, title, comparison):
    rect(slide, 0, 0, .14, 7.5, "4263EB")
    text(slide, .55, .30, 12.2, .50, title, 28 if len(title) < 50 else 23, bold=True)
    text(slide, .57, .89, 12.1, .45, f"A: {comparison.a.label}\nB: {comparison.b.label}", 11, "586B83")
    text(slide, .57, 1.43, 12.1, .30,
         f"Captured process {comparison.metric} | MiB | Exclusive grouping | Delta = B - A", 13, bold=True)


def footer(slide, comparison):
    line(slide, .55, 6.70, 12.75, 6.70)
    caution = ("Summed WS can double-count shared pages; not unique physical RAM."
               if comparison.use_ws else "Process private WS excludes shared pages and kernel allocations.")
    text(slide, .57, 6.80, 12.1, .57,
         caution + " Grouping does not measure total OS footprint.\n"
         "Remaining = Unclassified + Unknown + Ambiguous; configured exclusions are a policy, not proof of OS identity.\n"
         "Input records only: use wxemem --json --all. Unknown and overlapping matches are kept separate.",
         9.5, "586B83")


def abbreviated(name, length):
    return name if len(name) <= length else name[:length - 3] + "..."


def notes(comparison, rules, buckets, matches):
    rules_text = rules.text.replace("\r\n", "\n").replace("\r", "\n")
    result = [
        "Generated locally: " + datetime.now(timezone.utc).isoformat(),
        f"A: {comparison.a.path}\nSHA-256: {comparison.a.digest}",
        f"B: {comparison.b.path}\nSHA-256: {comparison.b.digest}",
        f"Groups: {rules.path}\nSHA-256: {rules.digest}\nRules snapshot (normalized line endings):\n{rules_text}",
        f"Requested metric={comparison.requested_metric}; effective metric={comparison.metric}.",
        f"Missing private WS: A={comparison.a.missing_private}; B={comparison.b.missing_private}.",
        "Substring matching is ASCII case-insensitive, as in wxememdiff. Multiple patterns within one group count once.",
        "Unlike wxememdiff's independent groups, an image matching multiple groups goes only into Ambiguous matches.",
        "Unmatched named images go into Unclassified; <unknown> always goes into Unknown images even if a rule matches it.",
        "Remaining includes Unclassified, Unknown and Ambiguous. It is not a measured OS footprint.",
        "Groups such as Tools can contain OS components. Exclusions describe the supplied policy, not intrinsic ownership.",
        "Total WS double-counts shared pages across processes; exclusive group assignment prevents a second kind of double counting, not shared-page duplication.",
        "If more than six groups are configured, five largest by max(A,B) are shown plus Other configured groups.",
        "All groups including zeros, rule matches, full names, instance counts and raw bytes follow.",
        f"A total={comparison.total_a}; B total={comparison.total_b}; delta={comparison.total_b - comparison.total_a}.",
    ]
    for bucket in buckets:
        result.append(f"GROUP {bucket.name}: A={bucket.a}, B={bucket.b}, delta={bucket.delta}")
        result.extend(json.dumps({"name": r.name, "matching_groups": matches[r.name],
                                  "a_bytes": r.a, "b_bytes": r.b, "delta_bytes": r.delta,
                                  "count_a": r.count_a, "count_b": r.count_b}, ensure_ascii=False)
                      for r in bucket.rows)
    return "\n".join(result)


def make_deck(comparison, rules, buckets, matches, title="WXE grouped process memory"):
    short_text(title, 60, "Title")
    visible = visible_buckets(buckets)
    if sum(v.a for v in visible) != comparison.total_a or sum(v.b for v in visible) != comparison.total_b:
        raise ValueError("Visible groups do not reconcile")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    slide = new_slide(prs)
    header(slide, title, comparison)
    maximum = max(comparison.total_a, comparison.total_b, MIB)
    scale = maximum * 1.05
    for letter, y, amount in (("A", 2.03, comparison.total_a), ("B", 2.87, comparison.total_b)):
        text(slide, .65, y, 1.15, .43, f"Report {letter}", 17, bold=True)
        rect(slide, 2.10, y, 8.65, .43, "E9EDF4")
        x = 2.10
        for i, bucket in enumerate(visible):
            value = bucket.a if letter == "A" else bucket.b
            width = value / scale * 8.65
            if value:
                rect(slide, x, y, width, .43, bucket.color, f"Composition-{letter}-{i}")
            x += width
        text(slide, 10.97, y, 1.80, .43, mib(amount), 19, bold=True,
             align=PP_ALIGN.RIGHT, name=f"Total-{letter}")
    line(slide, 2.10, 3.47, 10.75, 3.47)
    for fraction in (0, .25, .5, .75, 1):
        x = 2.10 + fraction * 8.65
        line(slide, x, 3.43, x, 3.51)
        text(slide, x - .46, 3.54, .92, .25, f"{fraction * scale / MIB:,.0f}",
             10, "586B83", align=PP_ALIGN.CENTER)
    for x, w, label in ((.93, 5.8, "Group"), (7.18, 1.40, "A (MiB)"),
                         (8.98, 1.40, "B (MiB)"), (10.77, 1.70, "B - A (MiB)")):
        text(slide, x, 3.94, w, .25, label, 11, "586B83", True,
             PP_ALIGN.LEFT if label == "Group" else PP_ALIGN.RIGHT)
    for i, bucket in enumerate(visible):
        y = 4.26 + .25 * i
        rect(slide, .65, y + .065, .14, .14, bucket.color)
        label = abbreviated(bucket.name, 42)
        if bucket.name == ROLLUP:
            label += f" ({len(bucket.members)})"
        text(slide, .93, y, 5.8, .23, label, 12, bold=not bucket.configured, name=f"Group-{i}")
        for x, w, value, field in ((7.18, 1.40, mib(bucket.a), "a"),
                                   (8.98, 1.40, mib(bucket.b), "b"),
                                   (10.77, 1.70, mib(bucket.delta, True), "delta")):
            text(slide, x, y, w, .23, value, min(12, 116 / len(value)), bold=not bucket.configured,
                 align=PP_ALIGN.RIGHT, name=f"Group-{i}-{field}")
    footer(slide, comparison)
    common_notes = notes(comparison, rules, buckets, matches)
    slide.notes_slide.notes_text_frame.text = common_notes

    slide = new_slide(prs)
    header(slide, title + " | Change", comparison)
    steps = waterfall_steps(comparison.total_a, visible, comparison.total_b)
    peak = max(max(s[2] for s in steps), MIB)
    bottom, height = 5.59, 3.46
    max_y = peak * 1.13
    width, left = 11.68 / len(steps), 1.04
    def plot_y(value):
        return bottom - value / max_y * height
    for fraction in (0, .25, .5, .75, 1):
        y = plot_y(max_y * fraction)
        line(slide, .95, y, 12.73, y)
        text(slide, .20, y - .12, .63, .24, f"{max_y * fraction / MIB:,.0f}",
             10, "586B83", align=PP_ALIGN.RIGHT)
    current = comparison.total_a
    for i, (name, lower, upper, value, color) in enumerate(steps):
        x, bar_width = left + i * width + width * .17, width * .66
        if upper > lower:
            rect(slide, x, plot_y(upper), bar_width, (upper - lower) / max_y * height,
                 color, f"Waterfall-{i}")
        else:
            line(slide, x, plot_y(upper), x + bar_width, plot_y(upper), color)
        amount = mib(value, 0 < i < len(steps) - 1)
        text(slide, left + i * width, plot_y(upper) - .32, width, .27,
             amount, min(12, 70 / len(amount)), bold=True, align=PP_ALIGN.CENTER,
             name=f"Waterfall-{i}-value")
        label = "\n".join(textwrap.wrap(abbreviated(name, 27), 14))
        text(slide, left + i * width, 5.73, width, .48, label,
             10, align=PP_ALIGN.CENTER, name=f"Waterfall-{i}-label")
        if i == 0:
            current = comparison.total_a
        elif i < len(steps) - 1:
            current += value
        if i < len(steps) - 1:
            next_x = left + (i + 1) * width + width * .17
            line(slide, x + bar_width, plot_y(current), next_x, plot_y(current), "98A9BC")
    remaining_a = sum(v.a for v in buckets if not v.configured)
    remaining_b = sum(v.b for v in buckets if not v.configured)
    text(slide, .65, 6.26, 12.0, .28,
         f"Remaining process {comparison.metric}: A {mib(remaining_a)}  |  B {mib(remaining_b)} MiB"
         f"     Net change (all groups): {mib(comparison.total_b - comparison.total_a, True)} MiB",
         12, "4263EB", True)
    footer(slide, comparison)
    slide.notes_slide.notes_text_frame.text = common_notes + "\nWaterfall steps:\n" + json.dumps(steps)
    prs.core_properties.title = title
    prs.core_properties.subject = f"A: {comparison.a.label}; B: {comparison.b.label}; grouping rules: {rules.path.name}"
    prs.core_properties.author = "wxemem-compare"
    buffer = BytesIO()
    prs.save(buffer)
    content = buffer.getvalue()
    verify_deck(content, comparison, visible, scale, steps, max_y, height)
    return content


def verify_deck(content, comparison, visible, scale, steps, max_y, plot_height):
    prs = Presentation(BytesIO(content))
    if len(prs.slides) != 2:
        raise ValueError("Expected composition and waterfall slides")
    for slide in prs.slides:
        for shape in slide.shapes:
            if (shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width
                    or shape.top + shape.height > prs.slide_height):
                raise ValueError(f"Shape exceeds slide bounds: {shape.name}")
    shapes = {s.name: s for s in prs.slides[0].shapes}
    for letter, total in (("A", comparison.total_a), ("B", comparison.total_b)):
        if shapes[f"Total-{letter}"].text != mib(total):
            raise ValueError("Composition total label mismatch")
        actual = sum(s.width for n, s in shapes.items() if n.startswith(f"Composition-{letter}-"))
        if abs(actual - Inches(total / scale * 8.65)) > len(visible):
            raise ValueError("Composition width does not match summed process memory")
    for i, bucket in enumerate(visible):
        for field, expected in (("a", mib(bucket.a)), ("b", mib(bucket.b)), ("delta", mib(bucket.delta, True))):
            if shapes[f"Group-{i}-{field}"].text != expected:
                raise ValueError("Group table value mismatch")
    shapes = {s.name: s for s in prs.slides[1].shapes}
    for i, (_, lower, upper, amount, _) in enumerate(steps):
        if shapes[f"Waterfall-{i}-value"].text != mib(amount, 0 < i < len(steps) - 1):
            raise ValueError("Waterfall value mismatch")
        if upper > lower:
            expected = Inches((upper - lower) / max_y * plot_height)
            if abs(shapes[f"Waterfall-{i}"].height - expected) > 1:
                raise ValueError("Waterfall height mismatch")


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--a", required=True, type=Path)
    parser.add_argument("--b", required=True, type=Path)
    parser.add_argument("--groups", required=True, type=Path, help="GroupN: name followed by substring patterns")
    parser.add_argument("--output", required=True, type=Path, help="New .pptx; never overwrites")
    parser.add_argument("--metric", choices=("ws", "private-ws", "auto"), default="ws")
    parser.add_argument("--label-a")
    parser.add_argument("--label-b")
    parser.add_argument("--title", default="WXE grouped process memory")
    args = parser.parse_args(argv)
    try:
        output = args.output.resolve()
        if output.suffix.lower() != ".pptx":
            raise ValueError("Output must have a .pptx extension")
        if output.exists():
            raise FileExistsError(f"Output already exists: {output}")
        a, b = load_snapshot(args.a, args.label_a), load_snapshot(args.b, args.label_b)
        if a.path.samefile(b.path):
            raise ValueError("A and B must be distinct report files")
        rules = load_groups(args.groups)
        comparison, buckets, matches = classify(a, b, rules, args.metric)
        if comparison.use_ws:
            print("CAUTION: Summed total WS can double-count shared pages; this is not total OS physical RAM.", file=sys.stderr)
        if comparison.requested_metric == "auto" and comparison.use_ws:
            print(f"WARNING: Private WS missing (A={a.missing_private}, B={b.missing_private}); both reports use WS.", file=sys.stderr)
        ambiguous = next(v for v in buckets if v.name == AMBIGUOUS)
        if ambiguous.rows:
            print(f"WARNING: {len(ambiguous.rows)} image names match multiple groups; assigned only to Ambiguous matches.", file=sys.stderr)
        content = make_deck(comparison, rules, buckets, matches, args.title)
        output.parent.mkdir(parents=True, exist_ok=True)
        with output.open("xb") as stream:
            try:
                stream.write(content)
            except OSError:
                stream.close()
                output.unlink()
                raise
        print(f"Created: {output}\nGroups: {rules.path}\nMetric: {comparison.metric}; Delta = B - A")
        for v in buckets:
            print(f"{v.name}: A={mib(v.a)} MiB; B={mib(v.b)} MiB; delta={mib(v.delta, True)} MiB")
        return 0
    except (OSError, ValueError) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
