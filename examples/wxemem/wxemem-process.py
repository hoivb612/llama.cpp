"""Compare captured process memory by image name; render four top-five PPTX tables."""

import argparse
from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from compare import MIB, infer_label, short_text, unique_object

TOP = 5
WS_CAUTION = "CAUTION: Summed WS can double-count shared pages; it is not unique physical RAM or a reconciliation of Used RAM."
GROUPS = (
    ("new", "NEW processes", "B only", "4263EB"),
    ("gone", "GONE processes", "A only", "12A594"),
    ("growers", "TOP GROWERS", "Present in both", "C77A20"),
    ("shrinkers", "TOP SHRINKERS", "Present in both", "25838A"),
)


@dataclass(frozen=True)
class Process:
    name: str
    ws: int
    private_ws: int | None


@dataclass(frozen=True)
class Snapshot:
    path: Path
    label: str
    digest: str
    processes: tuple[Process, ...]

    @property
    def missing_private(self):
        return sum(p.private_ws is None for p in self.processes)


@dataclass(frozen=True)
class Row:
    name: str
    a: int
    b: int
    count_a: int
    count_b: int

    @property
    def delta(self):
        return self.b - self.a


@dataclass(frozen=True)
class Comparison:
    a: Snapshot
    b: Snapshot
    use_ws: bool
    groups: dict[str, tuple[Row, ...]]
    total_a: int
    total_b: int
    requested_metric: str

    @property
    def metric(self):
        if self.use_ws:
            return "Total WS (fallback)" if self.requested_metric == "auto" else "Total WS"
        return "private WS"


def memory_value(value, location):
    if type(value) is not int or not 0 <= value <= 2**63 - 1:
        raise ValueError(f"{location} must be a nonnegative signed-64-bit byte count")
    return value


def load_snapshot(path, label=None):
    path = Path(path).resolve(strict=True)
    raw = path.read_bytes()
    data = json.loads(raw.decode("utf-8-sig"), object_pairs_hook=unique_object)
    if not isinstance(data, dict) or not isinstance(data.get("processes"), list):
        raise ValueError(f"{path}: report must contain a processes array")
    processes = []
    for i, p in enumerate(data["processes"]):
        where = f"{path}: processes[{i}]"
        if not isinstance(p, dict):
            raise ValueError(f"{where} must be an object")
        name = p.get("name")
        if not isinstance(name, str) or not name.strip() or any(ord(c) < 32 for c in name):
            raise ValueError(f"{where}.name must be a nonempty printable image name")
        ws = memory_value(p.get("ws_bytes"), f"{where}.ws_bytes")
        private = (memory_value(p["private_ws_bytes"], f"{where}.private_ws_bytes")
                   if "private_ws_bytes" in p else None)
        processes.append(Process(name, ws, private))
    return Snapshot(path, short_text(label if label is not None else infer_label(path), 90, "Report label"),
                    hashlib.sha256(raw).hexdigest(), tuple(processes))


def aggregate(snapshot, use_ws):
    result = {}
    for process in snapshot.processes:
        amount = process.ws if use_ws else process.private_ws
        if amount is None:
            raise ValueError("Private WS aggregation requires complete private_ws_bytes")
        total, count = result.get(process.name, (0, 0))
        result[process.name] = total + amount, count + 1
    return result


def compare_snapshots(a, b, metric="ws"):
    if metric not in ("ws", "private-ws", "auto"):
        raise ValueError(f"Unsupported process metric: {metric}")
    missing_private = a.missing_private > 0 or b.missing_private > 0
    if metric == "private-ws" and missing_private:
        raise ValueError("Requested private WS is unavailable for some records; use --metric ws or --metric auto")
    use_ws = metric == "ws" or (metric == "auto" and missing_private)
    aa, bb = aggregate(a, use_ws), aggregate(b, use_ws)
    groups = {key: [] for key, *_ in GROUPS}
    for name in sorted(aa.keys() | bb.keys()):
        amount_a, count_a = aa.get(name, (0, 0))
        amount_b, count_b = bb.get(name, (0, 0))
        row = Row(name, amount_a, amount_b, count_a, count_b)
        if not count_a:
            groups["new"].append(row)
        elif not count_b:
            groups["gone"].append(row)
        elif row.delta > 0:
            groups["growers"].append(row)
        elif row.delta < 0:
            groups["shrinkers"].append(row)
    groups["new"].sort(key=lambda r: (-r.b, r.name))
    groups["gone"].sort(key=lambda r: (-r.a, r.name))
    groups["growers"].sort(key=lambda r: (-r.delta, r.name))
    groups["shrinkers"].sort(key=lambda r: (r.delta, r.name))
    total_a, total_b = sum(v[0] for v in aa.values()), sum(v[0] for v in bb.values())
    if sum(r.delta for rows in groups.values() for r in rows) != total_b - total_a:
        raise ValueError("Process category deltas do not reconcile")
    return Comparison(a, b, use_ws, {k: tuple(v) for k, v in groups.items()}, total_a, total_b, metric)


def mib(value, signed=False):
    amount = abs(value) / MIB
    sign = "-" if value < 0 else "+" if signed else ""
    number = (f"{amount:,.2f}" if amount >= .01 or value == 0
              else f"{amount:.4f}" if amount >= .0001 else "<0.0001")
    return sign + number


def category_summary(key, rows):
    count = len(rows)
    procs = sum(r.count_b if key == "new" else r.count_a for r in rows)
    counts = f"{count} images" + (f", {procs} procs" if key in ("new", "gone") else "")
    return f"{counts} | total (all): {mib(sum(r.delta for r in rows), True)} MiB | shown {min(TOP, count)}"


def display_name(row, key):
    suffix = f" (x{row.count_b if key == 'new' else row.count_a})" if key in ("new", "gone") else ""
    limit = 39 - len(suffix)
    return (row.name if len(row.name) <= limit else row.name[:limit - 3] + "...") + suffix


def make_deck(comparison, title="WXE process memory comparisons"):
    short_text(title, 65, "Title")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("F7F9FC")

    def rect(x, y, w, h, color):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()

    def text(x, y, w, h, value, size=12, color="182A43", bold=False, right=False, name=None):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            shape.name = name
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, row in enumerate(value.split("\n")):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.text = row
            p.alignment = PP_ALIGN.RIGHT if right else PP_ALIGN.LEFT
            p.font.name, p.font.size, p.font.bold = "Aptos", Pt(size), bold
            p.font.color.rgb = RGBColor.from_string(color)
            p.space_after = Pt(0)
        return shape

    rect(0, 0, .14, 7.5, "4263EB")
    text(.55, .30, 12.2, .50, title, 29 if len(title) < 50 else 24, bold=True)
    text(.57, .89, 12.1, .45, f"A: {comparison.a.label}\nB: {comparison.b.label}", 11, "586B83")
    metric = f"By image name | {comparison.metric} | MiB | Delta = B - A"
    text(.57, 1.40, 8.1, .30, metric, 14, bold=True)
    text(8.70, 1.40, 4.0, .30,
         f"Net process delta: {mib(comparison.total_b - comparison.total_a, True)} MiB",
         13, "4263EB", True, True)
    for panel, (key, heading, membership, color) in enumerate(GROUPS):
        x = .55 if panel % 2 == 0 else 6.85
        y = 1.87 if panel < 2 else 4.27
        rect(x, y, 5.95, 2.28, "FFFFFF")
        rect(x, y, .045, 2.28, color)
        text(x + .17, y + .08, 3.4, .30, heading, 17, color, True)
        text(x + 3.4, y + .10, 2.37, .25, membership, 10, "586B83", right=True)
        rows = comparison.groups[key]
        text(x + .17, y + .43, 5.6, .25, category_summary(key, rows), 10, "586B83")
        text(x + .17, y + .76, 2.95, .23, "Image name" + (" (instances)" if key in ("new", "gone") else ""), 10, "586B83", True)
        for offset, label in ((3.18, "A"), (4.03, "B"), (4.88, "Delta")):
            text(x + offset, y + .76, .83, .23, label, 10, "586B83", True, True)
        if not rows:
            text(x + .17, y + 1.08, 5.4, .28, "None in captured input", 12, "586B83")
        for i, row in enumerate(rows[:TOP]):
            yy = y + 1.04 + i * .235
            value = display_name(row, key)
            text(x + .17, yy, 2.98, .225, value, min(12, 405 / len(value)),
                 name=f"{key}-row-{i}-name")
            for offset, label, number in ((3.18, "a", mib(row.a)),
                                          (4.03, "b", mib(row.b)),
                                          (4.88, "delta", mib(row.delta, True))):
                text(x + offset, yy, .83, .225, number, min(11, 74 / len(number)),
                     color if label == "delta" else "182A43", label == "delta", True,
                     name=f"{key}-row-{i}-{label}")
    rect(.55, 6.68, 12.2, .012, "DCE3ED")
    foot = (WS_CAUTION if comparison.use_ws else
            "Process private WS is not total OS Used RAM. Image names are matched exactly, as in wxememdiff.")
    foot += ("\nTop 5 per category; totals include all input rows. Use wxemem --json --all to avoid truncated process lists."
             "\nNew/Gone means snapshot membership, not installation. <unknown> aggregates may contain different processes.")
    text(.57, 6.78, 12.13, .59, foot, 10, "9A5916" if comparison.use_ws else "586B83")

    notes = [
        "Generated locally: " + datetime.now(timezone.utc).isoformat(),
        f"A: {comparison.a.path}\nSHA-256: {comparison.a.digest}",
        f"B: {comparison.b.path}\nSHA-256: {comparison.b.digest}",
        "Reference: wxememdiff.cpp loadSnapshot and User-mode process attribution sections.",
        "Exact, case-sensitive image-name aggregation; instance counts include all records for each name.",
        "Delta = B - A. This is intentionally opposite to the RAM-breakdown generator's A - B summary.",
        f"Requested metric: {comparison.requested_metric}; effective metric: {comparison.metric}; "
        f"missing private WS A={comparison.a.missing_private}, B={comparison.b.missing_private}.",
        "Default ws: use ws_bytes for all processes. Explicit private-ws requires complete private_ws_bytes. "
        "Auto: use private WS when complete, otherwise fall back to total WS for BOTH snapshots. Invalid supplied values are rejected.",
        WS_CAUTION,
        f"Sum of captured process {comparison.metric}: A={comparison.total_a} bytes; B={comparison.total_b} bytes; "
        f"B-A={comparison.total_b - comparison.total_a} bytes.",
        "New/Gone includes zero-memory images according to presence; unchanged common images are not ranked.",
        "Rank by byte magnitude before rounding. Equal values use image-name order for deterministic output.",
        "Category totals count all matching images and instances, not just displayed rows. Only up to five rows per category are shown.",
        "Inputs may be top-N captures. New/Gone does not prove a process was absent from the full system; use --all captures.",
        "Unknown image names are retained to match the source but do not identify a common executable.",
        "MiB = 2^20 bytes. C++ humanBytes uses binary scaling but labels it MB/KB. This slide uses explicit binary units.",
        "Full names and byte counts follow, including rows omitted from the slide:",
    ]
    for key, *_ in GROUPS:
        notes.append(key.upper())
        notes.extend(json.dumps({"name": r.name, "a_bytes": r.a, "b_bytes": r.b,
                                 "delta_bytes": r.delta, "count_a": r.count_a, "count_b": r.count_b},
                                ensure_ascii=False) for r in comparison.groups[key])
    slide.notes_slide.notes_text_frame.text = "\n".join(notes)
    prs.core_properties.title = title
    prs.core_properties.subject = f"A: {comparison.a.label}; B: {comparison.b.label}"
    prs.core_properties.author = "wxemem-compare"
    stream = BytesIO()
    prs.save(stream)
    content = stream.getvalue()
    verify_deck(content, comparison)
    return content


def verify_deck(content, comparison):
    prs = Presentation(BytesIO(content))
    if len(prs.slides) != 1:
        raise ValueError("Expected one process-comparison slide")
    slide = prs.slides[0]
    shapes = {s.name: s for s in slide.shapes}
    labels = "\n".join(s.text for s in slide.shapes if s.has_text_frame)
    if f"| {comparison.metric} |" not in labels:
        raise ValueError("Effective metric label missing")
    if comparison.use_ws and WS_CAUTION not in labels:
        raise ValueError("Total WS double-counting caution missing")
    for shape in slide.shapes:
        if (shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width
                or shape.top + shape.height > prs.slide_height):
            raise ValueError("Shape exceeds slide bounds")
    for key, *_ in GROUPS:
        rows = comparison.groups[key][:TOP]
        names = [s for s in shapes if s.startswith(key + "-row-") and s.endswith("-name")]
        if len(names) != len(rows):
            raise ValueError("Displayed row count mismatch")
        for i, row in enumerate(rows):
            for field, expected in (("name", display_name(row, key)), ("a", mib(row.a)),
                                    ("b", mib(row.b)), ("delta", mib(row.delta, True))):
                if shapes[f"{key}-row-{i}-{field}"].text != expected:
                    raise ValueError(f"Displayed value mismatch: {key}, row {i}, {field}")


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--a", required=True, type=Path)
    parser.add_argument("--b", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path, help="New .pptx file; never overwritten")
    parser.add_argument("--label-a")
    parser.add_argument("--label-b")
    parser.add_argument("--title", default="WXE process memory comparisons")
    parser.add_argument("--metric", choices=("ws", "private-ws", "auto"), default="ws",
                        help="ws (default): total WS; private-ws: require private WS; auto: private WS with symmetric WS fallback")
    args = parser.parse_args(argv)
    try:
        output = args.output.resolve()
        if output.suffix.lower() != ".pptx":
            raise ValueError("Output must have a .pptx extension")
        if output.exists():
            raise FileExistsError(f"Output already exists: {output}")
        a, b = load_snapshot(args.a, args.label_a), load_snapshot(args.b, args.label_b)
        if a.path.samefile(b.path):
            raise ValueError("A and B must be distinct report files")
        comparison = compare_snapshots(a, b, args.metric)
        if comparison.use_ws:
            print(WS_CAUTION, file=sys.stderr)
            if args.metric == "auto":
                print(f"WARNING: Missing private WS: A={a.missing_private}, B={b.missing_private}; "
                      "BOTH reports use total WS.", file=sys.stderr)
        content = make_deck(comparison, args.title)
        output.parent.mkdir(parents=True, exist_ok=True)
        with output.open("xb") as stream:
            try:
                stream.write(content)
            except OSError:
                stream.close()
                output.unlink()
                raise
        print(f"Created: {output}\nMetric: {comparison.metric}; Delta = B - A")
        print(f"Sum of captured process {comparison.metric}: A={mib(comparison.total_a)} MiB; "
              f"B={mib(comparison.total_b)} MiB; B-A={mib(comparison.total_b - comparison.total_a, True)} MiB")
        for key, heading, *_ in GROUPS:
            print(f"{heading}: {category_summary(key, comparison.groups[key])}")
        return 0
    except (OSError, ValueError) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
